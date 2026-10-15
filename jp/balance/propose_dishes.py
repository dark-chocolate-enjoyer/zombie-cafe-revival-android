#!/usr/bin/env python3
"""Dish balance PROPOSAL for the JP build (report-only; nothing is applied).

v2 (2026-10-08): every dish is treated the way the owner treated the most similar English
dishes going from EN vanilla to their v1.1 (EN_BALANCE_DIRECTION.md, DECISIONS D-06):

  1. Port       - the JP dish is identical to EN vanilla: copy the owner's EN v1.1 numbers.
  2. EN pattern - every other dish: find the K most similar free EN dishes (cook time, margin,
                  ingredient cost, $/serving, unlock level, $/min, all in vanilla terms), take
                  the one whose profit change is the median of the K ("modelled on"), and apply
                  the same change: its cost/servings multipliers and $/serving bump. If a flavour
                  number or the 900-servings cap blocks a lever, the same profit gain is reached
                  with the other levers. If the owner left those dishes alone, so does this.
  Cook time, XP and unlock level never change (EN held XP/min at vanilla); $/serving never drops;
  flavour numbers (years, repdigits like 555/777, 512/256/128...) are never edited.

  v1 used a single margin target per cook length; a leave-one-out test against the owner's real
  EN changes showed it missed the owner's pattern on 40 of 173 dishes (this method: ~21).

Outputs (report-only): proposals/dishes_proposal_v2.{csv,xlsx}, copied to
Desktop/ZCafeStuff/Japanese APK Restoration/Review/. Approved rows later become
jp/balance/overrides/*.csv rows (D-07).
"""
import csv
import json
import math
import os
import shutil
import statistics as st
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, os.path.join(HERE, '..', 'english', '.pylib'))
from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import Reference, ScatterChart, Series  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

SNAP = os.path.join(HERE, 'snapshots', 'foods_current.csv')
EN_REF = os.path.join(HERE, 'reference', 'en_food_vanilla_vs_v11.csv')
EN_V11_JSON = os.path.join(REPO, 'src', 'assets', 'data', 'foodData.bin.mid.json')
OUT = os.path.join(HERE, 'proposals')
REVIEW = '/mnt/c/Users/New/Desktop/ZCafeStuff/Japanese APK Restoration/Review'
STEM = 'dishes_proposal_v2'
DESKTOP_NAME = 'JP Dishes - Proposal v2 (2026-10-08)'

K = 5                     # similar EN dishes per JP dish (leave-one-out tested: 3-7 all similar)
BULK_SERVINGS = 900       # EN never pushed servings past 900
FILL_SERV_UP = 0.35       # top-up limits when a lever is blocked (EN capped serving lifts at +35%)
FILL_COST_CUT = 0.75
NOISE = 0.03              # changes under 3% are rounding noise, dropped
EN_MAX_PPS = 16           # no EN dish charges more per serving: above this there is nothing to compare with
EN_FREE_U8 = (0, 1, 2, 3, 4)  # EN level-track/raid/special dishes (not stove or Toxin-cookbook food)

FLAVOUR = {128, 256, 512, 214, 314, 1031, 1224, 1314, 1217, 1987, 2013, 2014, 2016}
UNLOCK = {1: 'Level track', 0: 'Raid recipe', 2: 'Raid recipe (also sold)', 3: 'Special',
          5: 'Event recipe', 4: 'Event recipe (late)'}


def is_flavour(n):
    s = str(n)
    return n in FLAVOUR or (len(s) >= 3 and len(set(s)) == 1) or (len(s) >= 4 and s.endswith('555'))


def step(n):
    return 1 if n < 20 else 5 if n < 200 else 10 if n < 1000 else 25 if n < 5000 else 50


def down(n):
    s = step(n)
    return max(1, int(n // s * s))


def up(n):
    s = step(n)
    return int(math.ceil(n / s) * s)


def serv_round(n, how=round):
    s = 5 if n < 1000 else 25
    return int(how(n / s) * s)


def band(lvl):
    return '0-5' if lvl <= 5 else '6-10' if lvl <= 10 else '11-15' if lvl <= 15 else '16-20' if lvl <= 20 \
        else '21-30' if lvl <= 30 else '31+'


def pct(a, b):
    return f'{(b - a) / a:+.0%}'


# ---------------------------------------------------------------- EN reference pool
def en_pool():
    """Free EN dishes with their vanilla state and the owner's v1.1 result."""
    v11 = json.load(open(EN_V11_JSON, encoding='utf-8'))
    pool = []
    for r, e in zip(csv.DictReader(open(EN_REF, encoding='utf-8-sig')), v11):
        if e['U8'] not in EN_FREE_U8:
            continue
        d = dict(name=r['name'], lvl=int(r['unlock_level']), c=int(r['cost_vanilla']), ck=int(r['cook_min_vanilla']),
                 s=int(r['servings_vanilla']), p=int(r['pps_vanilla']), c1=int(r['cost_v11']),
                 s1=int(r['servings_v11']), p1=int(r['pps_v11']))
        d['pv'], d['po'] = d['s'] * d['p'] - d['c'], d['s1'] * d['p1'] - d['c1']
        if d['pv'] > 0:
            d['mult'] = d['po'] / d['pv']
            d['f'] = feats(d['c'], d['ck'], d['s'], d['p'], d['lvl'])
            pool.append(d)
    return pool


def feats(c, ck, s, p, lvl):
    profit = s * p - c
    return (math.log2(max(ck, 1)), math.log2(max(profit / c, 0) + 0.1), math.log2(c), math.log2(p), lvl / 20,
            math.log2(max(profit / max(ck, 1), 0.05)))


def dist(f, d):
    return math.sqrt(sum((a - b) ** 2 for a, b in zip(f, d['f'])))


def similar(f, pool, skip=None):
    nb = sorted((d for d in pool if d is not skip), key=lambda d: dist(f, d))[:K]
    return nb, sorted(nb, key=lambda d: d['mult'])[K // 2]


def loose_threshold(pool):
    """95th percentile of how far EN dishes are from their own K nearest EN dishes."""
    ds = sorted(st.mean(dist(d['f'], n) for n in similar(d['f'], pool, skip=d)[0]) for d in pool)
    return ds[int(0.95 * (len(ds) - 1))]


def en_change_text(d):
    parts = [f'{lab} {pct(a, b)}' for lab, a, b in (('cost', d['c'], d['c1']), ('servings', d['s'], d['s1']))
             if a != b]
    if d['p1'] != d['p']:
        parts.append(f"$/serving {d['p']}->{d['p1']}")
    return ', '.join(parts) or 'unchanged'


def transfer(c, s, p, ex, flags):
    """Apply exemplar ex's EN change to a dish (cost c, servings s, $/serving p)."""
    target = (s * p - c) * ex['mult']
    c2, s2, p2 = c, s, p
    blocked = False
    if ex['c1'] != ex['c']:
        if is_flavour(c):
            blocked = True
        else:
            v = c * ex['c1'] / ex['c']
            c2 = down(v) if v < c else up(v)
    if ex['s1'] != ex['s']:
        v = s * ex['s1'] / ex['s']
        if is_flavour(s) or (v > s and s >= BULK_SERVINGS):
            blocked = True
        else:
            s2 = serv_round(min(v, max(s, BULK_SERVINGS)) if v > s else v)
    if ex['p1'] > ex['p']:
        p2 = max(p + 1, round(p * ex['p1'] / ex['p']))
    if ex['mult'] > 1 and s2 * p2 - c2 < 0.9 * target:     # a lever was blocked: reach the gain another way
        if not is_flavour(c):
            c2 = max(min(c2, up(s2 * p2 - target)), up(c * (1 - FILL_COST_CUT)))
        if s2 * p2 - c2 < 0.9 * target and not is_flavour(s) and s < BULK_SERVINGS:
            cap = serv_round(min(max(s, BULK_SERVINGS), s * (1 + FILL_SERV_UP)), math.floor)
            s2 = max(s2, min(serv_round((target + c2) / p2, math.ceil), cap))
        if blocked:
            flags.append('Flavour/bulk number kept; gain moved to another number')
        if s2 * p2 - c2 < 0.9 * target:
            flags.append('Partial: caps stop it reaching the EN gain')
    # drop rounding noise; never lower $/serving; never lower profit (EN never did)
    c2 = c if abs(c2 - c) / c < NOISE else c2
    s2 = s if abs(s2 - s) / s < NOISE else s2
    p2 = max(p2, p)
    if s2 * p2 - c2 < s * p - c and s2 < s:      # a copied "fewer servings" trim would cut profit: skip the trim
        s2 = s
    if s2 * p2 - c2 < s * p - c:
        return c, s, p
    return c2, s2, p2


# ---------------------------------------------------------------- proposal
def main():
    pool = en_pool()
    loose = loose_threshold(pool)
    rows = list(csv.DictReader(open(SNAP, encoding='utf-8-sig')))
    out = []
    for r in rows:
        I = lambda k: int(r[k])  # noqa: E731
        cost, cook, serv, pps, xp = I('cost_price'), I('cook_minutes'), I('servings'), I('price_per_serving'), I('xp')
        lvl, u8 = I('unlock_level'), I('U8')
        profit = serv * pps - cost
        ppm = profit / max(cook, 1)
        has_en = r['en_match'] != ''
        flags, modelled, near = [], '', ''
        n = dict(cost=cost, cook=cook, serv=serv, pps=pps, xp=xp)

        if r['jp_equals_en_vanilla'] == 'True':
            n = dict(cost=int(r['en_cost_v11']), cook=int(r['en_cook_min_v11']), serv=int(r['en_servings_v11']),
                     pps=int(r['en_pps_v11']), xp=int(r['en_xp_v11']))
            same = n == dict(cost=cost, cook=cook, serv=serv, pps=pps, xp=xp)
            kind = 'Port EN v1.1 (no change)' if same else 'Port EN v1.1'
            why = ('Same dish as EN vanilla; ' + ('you left it unchanged in EN v1.1' if same
                                                   else 'your EN v1.1 numbers copied as-is'))
            group = 'Ported from EN'
        elif pps > EN_MAX_PPS:
            kind = 'EN pattern: hold'
            why = (f'${pps:,} per serving is above anything in your EN game (max ${EN_MAX_PPS}), so no EN dish '
                   'is comparable; held for your call')
            flags.append('No comparable EN dish')
            group = 'EN dish, JP-repriced' if has_en else 'JP-only'
        else:
            f = feats(cost, cook, serv, pps, lvl)
            nb, ex = similar(f, pool)
            if st.mean(dist(f, d) for d in nb) > loose:
                flags.append('Loose match: no EN dish is very close')
            c2, s2, p2 = transfer(cost, serv, pps, ex, flags)
            n.update(cost=c2, serv=s2, pps=p2)
            changed = (c2, s2, p2) != (cost, serv, pps)
            kind = 'EN pattern: buff' if changed else 'EN pattern: hold'
            modelled = f"{ex['name']} (EN Lv{ex['lvl']}: {en_change_text(ex)})"
            near = ', '.join(d['name'] for d in nb)
            why = (f"Treated like {ex['name']}, the middle of the {K} most similar EN dishes: "
                   + (f"you changed it {en_change_text(ex)} (profit ×{ex['mult']:.2f})" if ex['mult'] != 1 or
                      en_change_text(ex) != 'unchanged' else 'you left it unchanged'))
            group = 'EN dish, JP-repriced' if has_en else 'JP-only'
            if n['pps'] != pps:
                flags.append('$/serving raised (also changes tips)')

        new_profit = n['serv'] * n['pps'] - n['cost']
        new_ppm = new_profit / max(n['cook'], 1)
        if ppm >= 5:
            flags.append(f'Quick-cash dish already ({ppm:.1f} $/min)')
        elif new_ppm >= 7:
            flags.append(f'Becomes a quick-cash dish ({new_ppm:.1f} $/min)')
        if u8 in (4, 5):
            flags.append('Event dish: offline unlock undecided (START_HERE §8)')
        if xp / max(cook, 1) >= 2.5:
            flags.append(f'XP outlier ({xp / max(cook, 1):.1f} XP/min)')

        cur = dict(cost=cost, cook=cook, serv=serv, pps=pps, xp=xp)
        changes = [f'{label} {cur[k]:,}->{n[k]:,} ({pct(cur[k], n[k])})' for k, label in (
            ('cost', 'Cost'), ('serv', 'Servings'), ('pps', '$/serving'), ('cook', 'Cook min'), ('xp', 'XP'))
            if cur[k] != n[k]]

        def en_text(sfx):
            if not has_en:
                return ''
            return (f"${int(r['en_cost_' + sfx]):,} · {r['en_cook_min_' + sfx]}m · "
                    f"{r['en_servings_' + sfx]}×${r['en_pps_' + sfx]} · {r['en_xp_' + sfx]}xp")

        out.append({
            '#': int(r['index']), 'Dish': r['name'], 'Lv': lvl, 'How unlocked': UNLOCK.get(u8, f'U8={u8}'),
            'EN twin': r['en_match'] or 'JP-only',
            'Same as EN vanilla': 'Yes' if r['jp_equals_en_vanilla'] == 'True' else ('No' if has_en else ''),
            'Proposal': kind, 'Buff': round(new_ppm / ppm - 1, 3) if ppm > 0 else '',
            'Owner verdict': '', 'Owner notes': '',
            'Cost': cost, 'Cook min': cook, 'Servings': serv, '$/serving': pps, 'XP': xp,
            'New cost': n['cost'], 'New cook min': n['cook'], 'New servings': n['serv'], 'New $/serving': n['pps'],
            'New XP': n['xp'],
            'Profit': profit, 'New profit': new_profit, '$/min': round(ppm, 2), 'New $/min': round(new_ppm, 2),
            'Margin': round(profit / cost, 2), 'New margin': round(new_profit / n['cost'], 2),
            'XP/min': round(xp / max(cook, 1), 3), 'New XP/min': round(n['xp'] / max(n['cook'], 1), 3),
            'What changes': '; '.join(changes) or '—', 'Why': why, 'Modelled on (EN)': modelled,
            'Similar EN dishes': near, 'Flags': '; '.join(flags),
            'EN vanilla': en_text('vanilla'), 'EN v1.1 (yours)': en_text('v11'),
            'EN $/min van→v1.1': (f"{float(r['en_profit_per_min_vanilla']):.2f} → "
                                  f"{float(r['en_profit_per_min_v11']):.2f}") if has_en else '',
            'JP name': r['jp_name'], '_band': band(lvl), '_group': group,
        })

    os.makedirs(OUT, exist_ok=True)
    cols = [k for k in out[0] if not k.startswith('_')]
    csv_path = os.path.join(OUT, STEM + '.csv')
    with open(csv_path, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, cols, extrasaction='ignore')
        w.writeheader()
        for o in out:
            w.writerow({**o, 'Buff': f"{o['Buff']:+.0%}" if o['Buff'] != '' else ''})
    xlsx_path = os.path.join(OUT, STEM + '.xlsx')
    write_xlsx(xlsx_path, out, cols, backtest(pool))
    print(f'{os.path.relpath(csv_path, HERE)}, {os.path.relpath(xlsx_path, HERE)}: {len(out)} dishes, '
          f"{sum(o['What changes'] != '—' for o in out)} change")
    for kind in sorted({o['Proposal'] for o in out}):
        print(f"  {kind}: {sum(o['Proposal'] == kind for o in out)}")
    if os.path.isdir(REVIEW):
        shutil.copyfile(csv_path, os.path.join(REVIEW, DESKTOP_NAME + '.csv'))
        shutil.copyfile(xlsx_path, os.path.join(REVIEW, DESKTOP_NAME + '.xlsx'))
        print(f'copied to {REVIEW}')


def backtest(pool):
    """Leave-one-out: predict each EN dish's v1.1 profit from the other EN dishes."""
    err, base, miss = [], [], 0
    for d in pool:
        _, ex = similar(d['f'], pool, skip=d)
        c2, s2, p2 = transfer(d['c'], d['s'], d['p'], ex, [])
        e = abs(math.log(max(s2 * p2 - c2, 1) / d['po']))
        err.append(e)
        base.append(abs(math.log(d['pv'] / d['po'])))
        miss += e > math.log(1.5)
    return dict(n=len(pool), err=st.mean(err), base=st.mean(base), miss=miss)


# ---------------------------------------------------------------- workbook
# Change size = how many times better (or worse) for the player: new/old - 1 when higher is better,
# old/new - 1 when lower is better (half the cost = +100%, same shade as double the servings).
STEPS = [0.02, 0.10, 0.20, 0.35, 0.50, 0.75, 1.00, 1.50, 2.50]


def _ramp(*stops, n=len(STEPS)):
    """n colours evenly spaced along the given hex colour stops."""
    rgb = [[int(c[i:i + 2], 16) for i in (0, 2, 4)] for c in stops]
    out = []
    for k in range(n):
        t = k / (n - 1) * (len(rgb) - 1)
        j = min(int(t), len(rgb) - 2)
        out.append(''.join(f'{round(x + (y - x) * (t - j)):02X}' for x, y in zip(rgb[j], rgb[j + 1])))
    return out


GREEN = _ramp('EEF8E9', '7FCF7A', '145A14')   # barely better -> hugely better
RED = _ramp('FDEDEC', 'F0857D', '8B0000')     # barely worse -> hugely worse
DARK_FROM = 6                       # shades from this index on get white text
FILL = {k: PatternFill('solid', fgColor=v) for k, v in {
    'Port EN v1.1': 'DDEBF7', 'Port EN v1.1 (no change)': 'EDF3FA', 'EN pattern: buff': 'E2EFDA',
    'EN pattern: hold': 'F2F2F2', 'head': '305496', 'group': 'D9E1F2', 'edit': '7F6000'}.items()}
WHITE_BOLD = Font(bold=True, color='FFFFFF')
THIN = Border(bottom=Side(style='thin', color='BFBFBF'))
GROUPS = [('Dish', 'Dish', 'Same as EN vanilla'), ('Your review', 'Proposal', 'Owner notes'),
          ('Current JP values', 'Cost', 'XP'), ('Proposed: type over these to try your own numbers', 'New cost',
                                                 'New XP'),
          ('Effect (New columns recalculate when you edit)', 'Profit', 'New XP/min'),
          ('Explanation', 'What changes', 'Flags'),
          ('Your English game (same dish)', 'EN vanilla', 'EN $/min van→v1.1'), ('JP original', 'JP name', 'JP name')]
WIDTH = {'#': 5, 'Dish': 26, 'Lv': 5, 'How unlocked': 17, 'EN twin': 20, 'Same as EN vanilla': 9,
         'Proposal': 19, 'Buff': 9, 'Owner verdict': 12, 'Owner notes': 24, 'What changes': 44, 'Why': 55,
         'Modelled on (EN)': 36, 'Similar EN dishes': 40, 'Flags': 36, 'EN vanilla': 27, 'EN v1.1 (yours)': 27,
         'EN $/min van→v1.1': 14, 'JP name': 18}


def shade_rules(ws, rng, new, old, higher_is_better=True):
    """Shade cells in rng by how much better/worse `new` is than `old` (cell refs for the first row)."""
    better = f'{new}/{old}-1' if higher_is_better else f'{old}/{new}-1'
    worse = f'{old}/{new}-1' if higher_is_better else f'{new}/{old}-1'
    for expr, colours in ((better, GREEN), (worse, RED)):
        for i in reversed(range(len(STEPS))):
            font = Font(bold=True, color='FFFFFF') if i >= DARK_FROM else Font(bold=i >= 3)
            fill = PatternFill(fill_type='solid', start_color=colours[i], end_color=colours[i])  # Excel reads end_color
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND(ISNUMBER({new}),ISNUMBER({old}),{new}<>{old},IFERROR({expr},0)>={STEPS[i]})'],
                fill=fill, font=font, stopIfTrue=True))


def write_xlsx(path, out, cols, bt):
    wb = Workbook()
    readme(wb.active, out, bt)
    ws = wb.create_sheet('Dishes')
    c = {name: i + 1 for i, name in enumerate(cols)}
    L = {name: get_column_letter(i) for name, i in c.items()}
    for title, a, b in GROUPS:
        ws.merge_cells(start_row=1, start_column=c[a], end_row=1, end_column=c[b])
        cell = ws.cell(1, c[a], title)
        cell.font, cell.fill, cell.alignment = Font(bold=True), FILL['group'], Alignment(horizontal='center')
    for name, i in c.items():
        cell = ws.cell(2, i, name)
        cell.font = WHITE_BOLD
        cell.fill = FILL['edit'] if name.startswith('New ') and name not in (
            'New profit', 'New $/min', 'New margin', 'New XP/min') else FILL['head']
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[get_column_letter(i)].width = WIDTH.get(name, 9)
    ws.row_dimensions[2].height = 32
    for ri, o in enumerate(out, start=3):
        for name, i in c.items():
            ws.cell(ri, i, o[name])
        nc, ncm, ns, npp, nx = (f'{L[k]}{ri}' for k in ('New cost', 'New cook min', 'New servings', 'New $/serving',
                                                         'New XP'))
        ws[f"{L['New profit']}{ri}"] = f'={ns}*{npp}-{nc}'
        ws[f"{L['New $/min']}{ri}"] = f"=ROUND({L['New profit']}{ri}/MAX({ncm},1),2)"
        ws[f"{L['New margin']}{ri}"] = f"=IF({nc}>0,ROUND({L['New profit']}{ri}/{nc},2),\"\")"
        ws[f"{L['New XP/min']}{ri}"] = f'=ROUND({nx}/MAX({ncm},1),3)'
        ws[f"{L['Buff']}{ri}"] = f"=IF({L['$/min']}{ri}>0,{L['New $/min']}{ri}/{L['$/min']}{ri}-1,\"\")"
        ws.cell(ri, c['Buff']).number_format = '+0%;-0%;0%'
        ws.cell(ri, c['Proposal']).fill = FILL.get(o['Proposal'], FILL['EN pattern: hold'])
        for k in ('What changes', 'Why', 'Modelled on (EN)', 'Similar EN dishes', 'Flags'):
            ws.cell(ri, c[k]).alignment = Alignment(wrap_text=True, vertical='top')
        for i in c.values():
            ws.cell(ri, i).border = THIN
        for k in ('Cost', 'New cost', 'Profit', 'New profit', 'Servings', 'New servings'):
            ws.cell(ri, c[k]).number_format = '#,##0'
    last = len(out) + 2
    # Buff and What changes: shaded by the change in earnings per minute
    for col in ('Buff', 'What changes'):
        shade_rules(ws, f'{L[col]}3:{L[col]}{last}', f"${L['New $/min']}3", f"${L['$/min']}3")
    # every changed number, typed or calculated: green = better for the player, red = worse
    for cur, new, higher in (('Cost', 'New cost', False), ('Cook min', 'New cook min', False),
                             ('Servings', 'New servings', True), ('$/serving', 'New $/serving', True),
                             ('XP', 'New XP', True), ('Profit', 'New profit', True), ('$/min', 'New $/min', True),
                             ('Margin', 'New margin', True), ('XP/min', 'New XP/min', True)):
        shade_rules(ws, f'{L[new]}3:{L[new]}{last}', f'{L[new]}3', f'${L[cur]}3', higher)
    dv = DataValidation(type='list', formula1='"Approve,Reject,Change (see notes)"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{L['Owner verdict']}3:{L['Owner verdict']}{last}")
    ws.freeze_panes = ws.cell(3, c['Lv'])
    ws.auto_filter.ref = f'A2:{get_column_letter(len(cols))}{last}'
    summary(wb.create_sheet('Summary'), out, bt)
    wb.save(path)


def readme(ws, out, bt):
    ws.title = 'Read me'
    count = lambda k: sum(o['Proposal'] == k for o in out)  # noqa: E731
    changed = sum(o['What changes'] != '—' for o in out)
    lines = [
        ('JP dish balance — proposal v2 (2026-10-08)', 'title'),
        ('Report only. Nothing has been applied to the game. You decide every row.', 'bold'),
        ('', None),
        ('How to review', 'h'),
        ('1. Open the "Dishes" sheet: one row per dish (234), in game-file order. Use the filter arrows to sort or '
         'filter (e.g. Proposal, How unlocked, or sort by Buff).', None),
        ('2. Colours: every proposed number that changed is shaded GREEN if it is better for the player and RED if '
         f'worse, in {len(STEPS)} shades from barely (2%) to huge (×3.5 or more), so the darker the bigger the '
         'change. That covers cost, servings, $/serving, cook time, XP, profit, $/min, margin and XP/min. "Buff" '
         'and "What changes" are shaded by the change in earnings per minute. Uncoloured = unchanged.', None),
        ('3. Type over any "New …" number (brown headers) to try your own value: colours, Buff and Effect columns '
         'update.', None),
        ('4. Set "Owner verdict" (Approve / Reject / Change) and add notes. Save the file.', None),
        ('5. Tell Claude "apply the approved dishes": approved rows become jp/balance/overrides/*.csv, then build, '
         'install on the A54, play-test.', None),
        ('', None),
        ('Colour key', 'h'),
        ('', 'legend'),
        ('', None),
        (f'What is proposed ({changed} of {len(out)} dishes change)', 'h'),
        (f'Port EN v1.1 — {count("Port EN v1.1")} dishes identical to EN vanilla get your EN v1.1 numbers exactly '
         f'(+{count("Port EN v1.1 (no change)")} you left unchanged in EN, so they stay).', None),
        (f'EN pattern — the other {count("EN pattern: buff") + count("EN pattern: hold")} dishes '
         f'({count("EN pattern: buff")} buffed, {count("EN pattern: hold")} held). Each one is matched to the '
         f'{K} most similar dishes of your English game (cook time, margin, ingredient cost, $/serving, unlock '
         'level, $/min) and gets the same change you gave the middle one of those: same % cost cut, same % '
         'servings change, same $/serving bump. If you left those dishes alone, this one is held too. The '
         '"Modelled on" and "Similar EN dishes" columns name them so you can check the match.', None),
        (f'Dishes charging more than ${EN_MAX_PPS} per serving (nothing in EN to compare with) are held for your call. '
         '"Loose match" flags dishes with no really close EN dish: check those matches first.', None),
        ('Cook time, XP and unlock level never change (you held XP per minute in EN). $/serving and profit never go '
         'down. Flavour numbers in JP event dishes (1987, 2014, 555, 777, 512/256/128 …) are kept; the same gain is '
         'moved to another number instead.', None),
        ('', None),
        ('Why this method (v2)', 'h'),
        (f'Tested on your own English game: predicting each of {bt["n"]} free EN dishes from the others, this '
         f'method reproduces your v1.1 profit change with a typical error of {math.exp(bt["err"]) - 1:.0%} '
         f'(misses by more than 1.5× on {bt["miss"]}). "Change nothing" is off by '
         f'{math.exp(bt["base"]) - 1:.0%}. v1 (one margin target per cook length) missed on 40 dishes: it '
         'buffed weak quick dishes you left alone and under-buffed the cheap dishes you made very profitable.',
         None),
        ('', None),
        ('Things to decide yourself', 'h'),
        ('• Event dishes (How unlocked = Event recipe) were online-event rewards; how they unlock offline is still '
         'open (START_HERE §8). They are matched against your FREE EN dishes. If they become paid/Toxin unlocks, '
         'your EN rule says premium food should earn clearly more.', None),
        ('• Quick-cash dishes and XP outliers are flagged. Some are held only because you held their EN '
         'equivalents.', None),
        ('', None),
        ('Column guide', 'h'),
        ('Margin = profit ÷ ingredient cost.  $/min = profit per cooking minute.  Buff = change in $/min.  '
         'EN columns: your English game\'s numbers for the same dish (cost · cook min · servings×$ · XP).', None),
        ('Source: jp/balance/propose_dishes.py (re-run to regenerate). Current values: '
         'jp/balance/snapshots/foods_current.csv.', None),
    ]
    row = 1
    for text, style in lines:
        if style == 'legend':
            row = legend(ws, row)
            continue
        cell = ws.cell(row, 1, text)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        if style == 'title':
            cell.font = Font(bold=True, size=14)
        elif style in ('h', 'bold'):
            cell.font = Font(bold=True, size=12 if style == 'h' else 11)
        row += 1
    ws.column_dimensions['A'].width = 120
    for i in range(len(STEPS)):
        ws.column_dimensions[get_column_letter(2 + i)].width = 8


def legend(ws, row):
    texts = [f'{t:+.0%}' if t < 1 else f'×{1 + t:g}' for t in STEPS]
    for colours, word in ((GREEN, 'better'), (RED, 'worse')):
        ws.cell(row, 1, f'{word.capitalize()} for the player (at least):').alignment = Alignment(horizontal='right')
        for i, (col, txt) in enumerate(zip(colours, texts)):
            cell = ws.cell(row, 2 + i, txt)
            cell.fill = PatternFill(fill_type='solid', start_color=col, end_color=col)
            cell.font = Font(bold=True, color='FFFFFF') if i >= DARK_FROM else Font(bold=i >= 3)
            cell.alignment = Alignment(horizontal='center')
        row += 1
    ws.cell(row, 1, 'Costs and cook time count by how much cheaper/faster: half the cost = ×2, the same shade as '
                    'double the servings.').alignment = Alignment(wrap_text=True)
    return row + 1


def summary(ws, out, bt):
    en = list(csv.DictReader(open(EN_REF, encoding='utf-8-sig')))
    med = lambda xs: round(st.median(xs), 2) if xs else ''  # noqa: E731

    def header(cells, height=30):
        ws.append(cells)
        r = ws.max_row
        for i in range(1, len(cells) + 1):
            cell = ws.cell(r, i)
            cell.font, cell.fill = WHITE_BOLD, FILL['head']
            cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
        ws.row_dimensions[r].height = height

    def title(text):
        ws.append([text])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=13)

    title('1. How the proposal moves JP, next to how you moved EN (vanilla → v1.1)')
    header(['Unlock levels', 'Dishes', 'Changed', 'Median $/min now', 'Median $/min proposed', 'Median margin now',
            'Median margin proposed', 'Median XP/min (unchanged)', 'EN: median $/min vanilla → v1.1',
            'EN: median margin vanilla → v1.1'], 45)
    bands = ('0-5', '6-10', '11-15', '16-20', '21-30', '31+')
    for b in bands:
        rows = [o for o in out if o['_band'] == b]
        e = [r for r in en if band(int(r['unlock_level'])) == b]
        em = lambda k: med([int(r['profit_' + k]) / int(r['cost_' + k]) for r in e])  # noqa: E731
        ws.append([b, len(rows), sum(o['What changes'] != '—' for o in rows), med([o['$/min'] for o in rows]),
                   med([o['New $/min'] for o in rows]), med([o['Margin'] for o in rows]),
                   med([o['New margin'] for o in rows]), med([o['XP/min'] for o in rows]),
                   f"{med([float(r['profit_per_min_vanilla']) for r in e])} → "
                   f"{med([float(r['profit_per_min_v11']) for r in e])}", f"{em('vanilla')} → {em('v11')}"])
    ws.append([])
    title('2. Do the JP-only dishes follow the same pattern as the dishes ported from your EN game?')
    header(['Unlock levels', 'Group', 'Dishes', 'Buffed', 'Median buff (buffed dishes)', 'Median $/min now',
            'Median $/min proposed', 'Median margin proposed'], 45)
    groups = ('Ported from EN', 'EN dish, JP-repriced', 'JP-only')
    for b in bands + ('All',):
        for g in groups:
            rows = [o for o in out if o['_group'] == g and (b == 'All' or o['_band'] == b)]
            if not rows:
                continue
            buffed = [o['Buff'] for o in rows if o['What changes'] != '—' and o['Buff'] != '']
            ws.append([b, g, len(rows), f'{len(buffed) / len(rows):.0%}', f'{st.median(buffed):+.0%}' if buffed
                       else '', med([o['$/min'] for o in rows]), med([o['New $/min'] for o in rows]),
                       med([o['New margin'] for o in rows])])
            if b == 'All':
                for i in range(1, 9):
                    ws.cell(ws.max_row, i).font = Font(bold=True)
    allm = {g: (med([o['$/min'] for o in out if o['_group'] == g]), med([o['New $/min'] for o in out if o['_group'] == g]))
            for g in groups}
    ws.append([f"Reading it: the buffs follow your EN pattern in every group, but JP-only dishes (mostly event dishes) "
               f"already earned {allm['JP-only'][0] / allm['Ported from EN'][0]:.1f}× more per minute than the "
               f"classic dishes before any change ({allm['JP-only'][0]} vs {allm['Ported from EN'][0]} $/min). "
               f"After the proposal the gap is {allm['JP-only'][1] / allm['Ported from EN'][1]:.1f}× "
               f"({allm['JP-only'][1]} vs {allm['Ported from EN'][1]}). Whether that is fine depends on how event "
               "dishes unlock offline (Read me)."])
    ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=False)
    ws.append([])
    title('3. By how the dish is unlocked')
    header(['How unlocked', 'Dishes', 'Changed', 'Median $/min now', 'Median $/min proposed', 'Median margin now',
            'Median margin proposed'])
    for u in UNLOCK.values():
        rows = [o for o in out if o['How unlocked'] == u]
        if rows:
            ws.append([u, len(rows), sum(o['What changes'] != '—' for o in rows), med([o['$/min'] for o in rows]),
                       med([o['New $/min'] for o in rows]), med([o['Margin'] for o in rows]),
                       med([o['New margin'] for o in rows])])
    ws.append([])
    ws.append([f'Method check: on your {bt["n"]} free EN dishes, this method predicts your v1.1 profit change '
               f'with a typical error of {math.exp(bt["err"]) - 1:.0%} (change-nothing: '
               f'{math.exp(bt["base"]) - 1:.0%}).'])
    for col, w in zip('ABCDEFGHIJ', (16, 22, 10, 10, 14, 14, 14, 14, 18, 18)):
        ws.column_dimensions[col].width = w
    chart_sheet(ws.parent.create_sheet('Chart data'), ws, out)


def chart_sheet(data, ws, out):
    """Proposed $/min by unlock level, one series per group (capped at 10 $/min for readability)."""
    data.append(['Group', 'Lv', 'Proposed $/min (capped at 10)', 'Dish'])
    spans = []
    for g in ('Ported from EN', 'EN dish, JP-repriced', 'JP-only'):
        start = data.max_row + 1
        for o in sorted((o for o in out if o['_group'] == g), key=lambda o: o['Lv']):
            data.append([g, o['Lv'], min(o['New $/min'], 10), o['Dish']])
        spans.append((g, start, data.max_row))
    ch = ScatterChart()
    ch.title = 'Proposed earnings per minute by unlock level'
    ch.x_axis.title, ch.y_axis.title = 'Unlock level', '$/min (capped at 10)'
    ch.style, ch.height, ch.width = 13, 11, 24
    for g, a, b in spans:
        s = Series(Reference(data, min_col=3, min_row=a, max_row=b), Reference(data, min_col=2, min_row=a, max_row=b),
                   title=g)
        s.marker.symbol, s.marker.size, s.graphicalProperties.line.noFill = 'circle', 6, True
        ch.series.append(s)
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 10
    ch.x_axis.scaling.min, ch.x_axis.scaling.max = 0, 70
    ch.x_axis.delete = ch.y_axis.delete = False
    ws.add_chart(ch, 'L2')


if __name__ == '__main__':
    main()
